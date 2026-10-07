from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Paleta de la app
AZUL = '206BC4'
GRIS_CLARO = 'F5F5F5'
GRIS_BORDE = 'D0D0D0'
ROJO = 'D63939'
AMARILLO = 'F59F00'
VERDE = '2FB344'


def _estilo_header(cell):
    cell.font = Font(bold=True, color='FFFFFF', size=11)
    cell.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')


def _borde_fino():
    lado = Side(border_style='thin', color=GRIS_BORDE)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _aplicar_bordes_y_anchos(ws, num_cols, ancho_minimo=12):
    borde = _borde_fino()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=num_cols):
        for cell in row:
            cell.border = borde

    for col_idx in range(1, num_cols + 1):
        max_len = ancho_minimo
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)) + 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 50)


def exportar_pedido(pedido):
    """
    Genera un Excel con la Orden de Compra de un pedido.

    Retorna: BytesIO con el archivo listo para descargar.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f'Pedido {pedido.pk}'

    detalles = pedido.detalles.select_related('producto').all()

    # ============================================
    # HEADER - Info del pedido
    # ============================================
    ws.merge_cells('A1:F1')
    ws['A1'] = 'ORDEN DE COMPRA'
    ws['A1'].font = Font(bold=True, size=18, color=AZUL)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:F2')
    ws['A2'] = f'Pedido #{pedido.pk}'
    ws['A2'].font = Font(bold=True, size=14)
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 22

    # ============================================
    # INFO GENERAL
    # ============================================
    ws['A4'] = 'Tienda:'
    ws['A4'].font = Font(bold=True)
    ws['B4'] = pedido.tienda.nombre

    ws['D4'] = 'Fecha:'
    ws['D4'].font = Font(bold=True)
    ws['E4'] = pedido.fecha_creacion.strftime('%d/%m/%Y %H:%M')

    ws['A5'] = 'Solicitado por:'
    ws['A5'].font = Font(bold=True)
    ws['B5'] = pedido.usuario.get_full_name() or pedido.usuario.username

    ws['D5'] = 'Estado:'
    ws['D5'].font = Font(bold=True)
    ws['E5'] = pedido.get_estatus_display()

    # ============================================
    # TABLA DE PRODUCTOS (empieza en fila 7)
    # ============================================
    headers = [
        'SKU',
        'Producto',
        'Cuadrante',
        'Cant. sugerida',
        'Cant. ajustada',
        'Cant. recibida',
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col, value=header)
        _estilo_header(cell)
    ws.row_dimensions[7].height = 22

    # Rellenar filas
    fila = 8
    total_sugerida = 0
    total_ajustada = 0
    total_recibida = 0

    for detalle in detalles:
        producto = detalle.producto
        cant_ajustada = detalle.cant_ajustada if detalle.cant_ajustada is not None else detalle.cant_sugerida
        cant_recibida = detalle.cant_recibida if detalle.cant_recibida is not None else 0

        ws.cell(row=fila, column=1, value=producto.sku)
        ws.cell(row=fila, column=2, value=producto.nombre)
        ws.cell(row=fila, column=3, value=producto.cuadrante or '—')
        ws.cell(row=fila, column=4, value=detalle.cant_sugerida)
        ws.cell(row=fila, column=5, value=cant_ajustada)
        ws.cell(row=fila, column=6, value=cant_recibida if detalle.cant_recibida is not None else '—')

        # Resaltar si hubo ajuste
        if detalle.cant_ajustada is not None and detalle.cant_ajustada != detalle.cant_sugerida:
            ws.cell(row=fila, column=5).fill = PatternFill(
                start_color=AMARILLO, end_color=AMARILLO, fill_type='solid'
            )

        total_sugerida += detalle.cant_sugerida
        total_ajustada += cant_ajustada
        total_recibida += cant_recibida
        fila += 1

    # Fila de totales
    ws.cell(row=fila, column=1, value='TOTAL').font = Font(bold=True)
    ws.cell(row=fila, column=4, value=total_sugerida).font = Font(bold=True)
    ws.cell(row=fila, column=5, value=total_ajustada).font = Font(bold=True)
    ws.cell(row=fila, column=6, value=total_recibida if total_recibida > 0 else '—').font = Font(bold=True)
    for col in range(1, 7):
        ws.cell(row=fila, column=col).fill = PatternFill(
            start_color=GRIS_CLARO, end_color=GRIS_CLARO, fill_type='solid'
        )

    # ============================================
    # ESTILOS Y ANCHOS
    # ============================================
    _aplicar_bordes_y_anchos(ws, num_cols=6)

    # Alinear números a la derecha
    for row in ws.iter_rows(min_row=8, max_row=ws.max_row, min_col=4, max_col=6):
        for cell in row:
            cell.alignment = Alignment(horizontal='right')

    # ============================================
    # FOOTER
    # ============================================
    ws.cell(row=fila + 2, column=1, value='Generado por el sistema de gestión - Granja').font = Font(
        italic=True, size=9, color='888888'
    )
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def exportar_lista_pedidos(pedidos):
    """
    Genera un Excel con una lista de pedidos.

    pedidos: queryset de Pedido
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pedidos'

    # Header
    ws.merge_cells('A1:H1')
    ws['A1'] = 'LISTA DE PEDIDOS'
    ws['A1'].font = Font(bold=True, size=16, color=AZUL)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    # Cabeceras
    headers = ['Pedido #', 'Fecha', 'Tienda', 'Solicitante', 'Estado', '# Productos', 'Total sugerido', 'Total recibido']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        _estilo_header(cell)
    ws.row_dimensions[3].height = 20

    # Datos
    fila = 4
    for pedido in pedidos:
        detalles = list(pedido.detalles.all())
        total_sugerido = sum(d.cant_sugerida for d in detalles)
        total_recibido = sum((d.cant_recibida or 0) for d in detalles)

        ws.cell(row=fila, column=1, value=f'#{pedido.pk}')
        ws.cell(row=fila, column=2, value=pedido.fecha_creacion.strftime('%d/%m/%Y'))
        ws.cell(row=fila, column=3, value=pedido.tienda.nombre)
        ws.cell(row=fila, column=4, value=pedido.usuario.get_full_name() or pedido.usuario.username)
        ws.cell(row=fila, column=5, value=pedido.get_estatus_display())
        ws.cell(row=fila, column=6, value=len(detalles))
        ws.cell(row=fila, column=7, value=total_sugerido)
        ws.cell(row=fila, column=8, value=total_recibido if total_recibido > 0 else '—')

        fila += 1

    _aplicar_bordes_y_anchos(ws, num_cols=8)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

